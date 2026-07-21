## extract_word.py
# import
import re

from docx import Document

from src.model_classes_parsing.base_classes_parsing import (
    Bullet,
    BulletList,
    Heading,
    HeadingMeta,
    TextBlock,
)

"""
ZIEL:
Für dein Projekt würde ich sogar einen Schritt weitergehen

Da dein Ziel ein GMP-Dokumentenparser ist und nicht nur ein PDF-Parser, würde ich die Extraktion künftig nach semantischen Dokumentelementen ausrichten, nicht nach Dateiformaten.

Das hieße:

PDF
        \
         \
DOCX ------> DocumentElement[]
         /
HTML ----/

mit gemeinsamen Klassen wie

Heading
Paragraph
BulletList
Table
Image
CodeBlock (optional)
"""


def extract_docx(f_path: str):
    # doc = read_docx()
    doc = Document(f_path)

    para_all = []
    container_id = 0

    for idx, p in enumerate(doc.paragraphs):
        # --> Heading, Paragraph / Normal, List_Bullet / Aufzählung,
        # List_Number, Table
        # st.write(f"Paragraph #{idx} (type = {p.style.name}):\n", p)
        # print()

        if idx < len(doc.paragraphs) - 1:
            bullet_next = doc.paragraphs[idx + 1].style.name == "Aufzählung"

        bullets = []

        match p.style.name:
            case "Heading 1" | "Heading 2":
                para_all.append(
                    Heading(
                        meta=HeadingMeta(level=re.findall(r"\d+", p.style.name)[0]),
                        text=p.text,
                        container_id=container_id,
                    )
                )
                container_id += 1

            case "Normal":
                para_all.append(
                    TextBlock(
                        # meta = TextBlockMeta(
                        #             level = re.findall(
                        #                             r'\d+',
                        #                             p.style.name
                        #                             )[0]
                        #             ),
                        text=p.text,
                        container_id=container_id,
                    )
                )
                container_id += 1

            case "Aufzählung":
                bullets.append(Bullet(text=p.text))

                if bullet_next is False:
                    para_all.append(
                        BulletList(
                            elements=bullets,
                            text="- " + "\n- ".join([b.text for b in bullets]),
                            container_id=container_id,
                        )
                    )

                    container_id += 1
                    bullets = []

    return para_all
