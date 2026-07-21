## extract_excel_helper.py
# import
import pandas as pd


def extract_xls_data(path):
    xls = pd.ExcelFile(path)

    xls_tbl = []
    for sheet in xls.sheet_names:
        xls_tbl.append(pd.read_excel(path, sheet_name=sheet))

    return xls_tbl


"""
from openpyxl import Workbook
from openpyxl.styles import Font
wb = Workbook()
ws = wb.active
ws["A1"] = "Product"
ws["B1"] = "Sales"
headers = ["A1", "B1"]
for cell in headers:
    ws[cell].font = Font(bold=True)
data = [
    ["Laptop", 1500],
    ["Mouse", 300],
    ["Keyboard", 500],
    ["Monitor", 1200]
]
for row in data:
    ws.append(row)
wb.save("sales_report.xlsx")
print("Report generated")
"""

"""
from openpyxl import load_workbook

workbook = load_workbook("sales.xlsx")
sheet = workbook.active

new_data = [100, 200, 150]  # Daily sales
sheet.append(new_data)

workbook.save("sales.xlsx")
print("Excel updated!")
"""


def extract_xls_format(path):
    """
    -> merged cells
    -> comments
    -> styles
    -> formulas
    -> positions
    """

    return


from openpyxl import load_workbook


def update_excel(path):
    wb = load_workbook(path)
    ws = wb.active

    ws["A1"] = "Processed Automatically"
    ws.append(["John Doe", 500, "Completed"])

    wb.save("updated.xlsx")


update_excel("template.xlsx")
