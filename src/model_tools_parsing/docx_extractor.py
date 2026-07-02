## docx_extractor.py
# import
from dataclasses import dataclass
from docx import Document
from pathlib import Path
import re
# import streamlit as st
from returns.result import Failure, Result, Success


from src.model_tools_parsing.base_extractor import BaseExtractor
from src.model_classes_parsing.base_classes_parsing import (
                                    Bullet, 
                                    BulletList,
                                    DocumentExtract,
                                    Heading,
                                    HeadingMeta,
                                    RawDocument,
                                    TextBlock
                                    )
from src.utils.dict_helper import save_dict



@dataclass
class DOCXCleanExtractor(BaseExtractor):

    def extract(
            self, 
            f_path: str
            ) -> Result[DocumentExtract, str]:
        doc = Document(f_path)

        para_clean = self._extract_per_paragraph(doc)
        
        text_clean = "\n\n".join([p.text for p in para_clean])
        doc_info = self.enricher._add_basic_metadata(
                                            text_clean, 
                                            para_clean, 
                                            RawDocument()
                                            )
        
        doc_info = self._add_head_context(doc_info)

        f_path = Path(f"{self.save_folder}/{self.save_name}")
        save_dict(data=doc_info.model_dump(
                            mode="json",
                            serialize_as_any=True,
                            ), 
                path=f_path)

        return Success(
                DocumentExtract(
                        doc_type="docx",
                        doc_name=self.doc_name,
                        text=doc_info.text,
                        elements=doc_info.elements,
                        meta=doc_info.meta,
                        )
                )


    def _add_head_context(self, doc_info: RawDocument) -> RawDocument:
        
        para_sorted = sorted(doc_info.elements, 
                             key=lambda e: e.container_id)

        text = []
        current_context = []
        blocks = []

        for para in para_sorted:
            c_type = para.container_type
            content = para.text or ""
            # ).strip()

            if not content:
                continue

            if c_type == "heading":
                level = para.meta.level or 1

                # alle Headings gleicher oder tieferer Ebene entfernen
                current_context = current_context[:level - 1]
                current_context.append(content)

                prefix = "#" * level
                text.append(f"{prefix} {content}")

            elif c_type in ["text_block", "bullet_list"]:
                para.context = current_context.copy()

                if c_type == "bullet_list":
                    text.append(content)
                else:
                    text.append(content)

                blocks.append(para)

        para.text = "\n\n".join(text)

        return doc_info 
    # {
    #         "text": "\n\n".join(text),
    #         "blocks": blocks,
    #     }

        # headings = []
        # non_heads = []
        # for para in doc_info: 
        #     if para.container_type == "heading":
        #         headings.append(para)

        #     else: 
        #         non_heads.append(para)

    
        # heads_sorted = sorted(headings, 
        #                      key=lambda p: p.container_id)
        # non_heads_sorted = sorted(non_heads, 
        #                      key=lambda p: p.container_id)
        
        # text = []
        # idx = 0
        # relevant_heads = []
        # head_lvl = None

        # for idx, para in enumerate(para_sorted):
        #     if para.container_type == "heading":
        #         if not head_lvl: 
        #             # or para.meta.level <= head_lvl:
        #             relevant_heads = []
                
        #         relevant_heads[head_lvl]

                
        #     else:
        #         para.context = relevant_heads

            

        # return 
    

    def _extract_per_paragraph(self, doc: Document) -> list:
    
        para_all = []
        container_id = 0

        for idx, p in enumerate(doc.paragraphs):
            if idx < len(doc.paragraphs) - 1:
                bullet_next = (doc.paragraphs[idx+1].style.name == "Aufzählung")

            bullets = []

            match p.style.name:
                case "Heading 1" | "Heading 2":
                    para_all.append(
                            Heading(
                                meta = HeadingMeta(
                                            level=re.findall(
                                                        r'\d+', 
                                                        p.style.name
                                                        )[0]
                                            ),
                                text = p.text,
                                container_id = idx
                                )
                            )
                    container_id += 1

                case "Normal":
                    para_all.append(
                            TextBlock(
                                # meta = TextBlockMeta(
                                #             level = re.findall(
                                #                             r'\d+', 
                                #                             p.style.name
                                #                             )[0]
                                #             ),
                                text = p.text,
                                container_id = idx
                                )
                            )
                    container_id += 1

                case "Aufzählung":

                    bullets.append(
                            Bullet(text = p.text)
                            )

                    if bullet_next is False:
                        para_all.append(
                                BulletList(
                                    elements = bullets,
                                    text = "- " + "\n- ".join([b.text for b in bullets]),
                                    container_id = idx
                                    )
                                )
                        
                        container_id += 1
                        bullets = []      

        return para_all